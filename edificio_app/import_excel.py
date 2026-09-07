"""
Importa los datos de tu Excel actual (LIBRO DIARIO, LECTURA DE AGUA, Referencias)
a la base de datos del sistema, para no empezar de cero.

Uso:
    python import_excel.py "Libro_I_O_Eva_1_-__2025.xlsx"
"""
import sys
from datetime import date

import openpyxl

from app import app
from models import db, Residente, MovimientoCaja, LecturaAgua, Configuracion


def get_or_create_residente(nombre, medidor=None):
    nombre = (nombre or '').strip().upper()
    if not nombre:
        return None
    r = Residente.query.filter_by(nombre=nombre).first()
    if not r:
        r = Residente(nombre=nombre, medidor=(medidor or '').strip())
        db.session.add(r)
        db.session.flush()
    return r


def importar(ruta_excel):
    wb = openpyxl.load_workbook(ruta_excel, data_only=True)

    with app.app_context():
        db.create_all()

        # --- Referencias -> Configuracion ---
        if 'Referencias' in wb.sheetnames:
            ws = wb['Referencias']
            cfg = Configuracion.query.first() or Configuracion()
            for row in ws.iter_rows(values_only=True):
                if not row or not row[0]:
                    continue
                etiqueta = str(row[0]).strip().upper()
                valor = row[1]
                if 'EXPENSA' in etiqueta and isinstance(valor, (int, float)):
                    cfg.monto_expensa = valor
                elif 'AGUA POR M3' in etiqueta and isinstance(valor, (int, float)):
                    cfg.precio_m3 = valor
                elif 'LIMPIEZA EDIFICIO' in etiqueta and isinstance(valor, (int, float)):
                    cfg.precio_limpieza_edificio = valor
                elif 'LIMPIEZA PARQUEO' in etiqueta and isinstance(valor, (int, float)):
                    cfg.precio_limpieza_parqueo = valor
            if not cfg.fecha_inicio_administracion:
                cfg.fecha_inicio_administracion = date.today()
            db.session.add(cfg)
            db.session.commit()
            print('Configuración importada.')

        # --- LECTURA DE AGUA -> Residentes + LecturaAgua ---
        if 'LECTURA DE AGUA 2025-2026' in wb.sheetnames:
            ws = wb['LECTURA DE AGUA 2025-2026']
            filas = list(ws.iter_rows(min_row=2, values_only=True))
            count = 0
            for row in filas:
                if not row or not row[4]:  # NOMBRE vacío
                    continue
                (_, fecha_lectura, mes, medidor, nombre, _fecha_multa,
                 lect_ant, lect_act, consumo, costo, expensas, multa,
                 total, pagado, obs) = (list(row) + [None] * 15)[:15]

                residente = get_or_create_residente(nombre, medidor)
                if not residente:
                    continue

                anio = fecha_lectura.year if hasattr(fecha_lectura, 'year') else date.today().year
                lectura = LecturaAgua(
                    residente_id=residente.id,
                    mes=str(mes).strip().upper() if mes else 'S/D',
                    anio=anio,
                    fecha_lectura=fecha_lectura.date() if hasattr(fecha_lectura, 'date') else None,
                    lectura_anterior=lect_ant or 0,
                    lectura_actual=lect_act or 0,
                    consumo_m3=consumo or 0,
                    costo_agua=costo or 0,
                    monto_expensa=expensas or 0,
                    multa=multa or 0,
                    total=total or 0,
                    pagado=bool(pagado),
                    observaciones=str(obs) if obs else None)
                db.session.add(lectura)
                count += 1
            db.session.commit()
            print(f'{count} lecturas de agua importadas.')

        # --- LIBRO DIARIO -> MovimientoCaja ---
        if 'LIBRO DIARIO' in wb.sheetnames:
            ws = wb['LIBRO DIARIO']
            filas = list(ws.iter_rows(min_row=2, values_only=True))
            count = 0
            for row in filas:
                if not row or not row[1] or not row[2]:
                    continue
                _, fecha, denom, detalle, debe, haber = (list(row) + [None] * 6)[:6]
                denom = str(denom).strip().upper() if denom else ''
                if denom not in ('INGRESOS', 'EGRESOS'):
                    continue
                monto = debe if denom == 'INGRESOS' else haber
                if not monto:
                    continue
                mov = MovimientoCaja(
                    fecha=fecha.date() if hasattr(fecha, 'date') else date.today(),
                    tipo='INGRESO' if denom == 'INGRESOS' else 'EGRESO',
                    detalle=str(detalle).strip() if detalle else '',
                    monto=float(monto))
                db.session.add(mov)
                count += 1
            db.session.commit()
            print(f'{count} movimientos de caja importados.')

    print('Importación completa.')


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print('Uso: python import_excel.py "ruta_al_archivo.xlsx"')
        sys.exit(1)
    importar(sys.argv[1])
