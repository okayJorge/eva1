import os
from datetime import date, datetime, timedelta
from io import BytesIO

from flask import (Flask, render_template, redirect, url_for, request,
                    flash, send_file, abort)
from flask_login import (LoginManager, login_user, logout_user,
                          login_required, current_user)
from werkzeug.security import generate_password_hash, check_password_hash
from sqlalchemy import func, extract
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from models import (db, Usuario, Residente, Configuracion, MovimientoCaja,
                     LecturaAgua, DeudaExtra)

BASE_DIR = os.path.abspath(os.path.dirname(__file__))

app = Flask(__name__)
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'cambia-esta-clave-en-produccion')
app.config['SQLALCHEMY_DATABASE_URI'] = os.environ.get(
    'DATABASE_URL', f"sqlite:///{os.path.join(BASE_DIR, 'edificio.db')}"
).replace('postgres://', 'postgresql://', 1)
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db.init_app(app)

login_manager = LoginManager()
login_manager.login_view = 'login'
login_manager.login_message = 'Debes iniciar sesión para continuar.'
login_manager.init_app(app)

MESES = ['ENERO', 'FEBRERO', 'MARZO', 'ABRIL', 'MAYO', 'JUNIO', 'JULIO',
         'AGOSTO', 'SEPTIEMBRE', 'OCTUBRE', 'NOVIEMBRE', 'DICIEMBRE']


@login_manager.user_loader
def load_user(user_id):
    return Usuario.query.get(int(user_id))


def get_config():
    cfg = Configuracion.query.first()
    if not cfg:
        cfg = Configuracion(fecha_inicio_administracion=date.today())
        db.session.add(cfg)
        db.session.commit()
    return cfg


# ---------------------------------------------------------------- AUTH ----

@app.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated:
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        username = request.form['username'].strip()
        password = request.form['password']
        user = Usuario.query.filter_by(username=username).first()
        if user and check_password_hash(user.password_hash, password):
            login_user(user)
            return redirect(url_for('dashboard'))
        flash('Usuario o contraseña incorrectos.', 'danger')
    return render_template('login.html')


@app.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('login'))


@app.route('/cambiar-clave', methods=['GET', 'POST'])
@login_required
def cambiar_clave():
    if request.method == 'POST':
        actual = request.form['actual']
        nueva = request.form['nueva']
        if not check_password_hash(current_user.password_hash, actual):
            flash('La contraseña actual no es correcta.', 'danger')
        elif len(nueva) < 6:
            flash('La nueva contraseña debe tener al menos 6 caracteres.', 'warning')
        else:
            current_user.password_hash = generate_password_hash(nueva)
            db.session.commit()
            flash('Contraseña actualizada.', 'success')
            return redirect(url_for('dashboard'))
    return render_template('cambiar_clave.html')


# ----------------------------------------------------------- DASHBOARD ----

@app.route('/')
@login_required
def dashboard():
    cfg = get_config()
    hoy = date.today()

    total_ingresos = db.session.query(func.coalesce(func.sum(MovimientoCaja.monto), 0)) \
        .filter(MovimientoCaja.tipo == 'INGRESO').scalar()
    total_egresos = db.session.query(func.coalesce(func.sum(MovimientoCaja.monto), 0)) \
        .filter(MovimientoCaja.tipo == 'EGRESO').scalar()
    saldo_caja = total_ingresos - total_egresos

    ingresos_mes = db.session.query(func.coalesce(func.sum(MovimientoCaja.monto), 0)) \
        .filter(MovimientoCaja.tipo == 'INGRESO',
                extract('month', MovimientoCaja.fecha) == hoy.month,
                extract('year', MovimientoCaja.fecha) == hoy.year).scalar()
    egresos_mes = db.session.query(func.coalesce(func.sum(MovimientoCaja.monto), 0)) \
        .filter(MovimientoCaja.tipo == 'EGRESO',
                extract('month', MovimientoCaja.fecha) == hoy.month,
                extract('year', MovimientoCaja.fecha) == hoy.year).scalar()

    deuda_lecturas = db.session.query(func.coalesce(func.sum(LecturaAgua.total), 0)) \
        .filter(LecturaAgua.pagado.is_(False)).scalar()
    deuda_extra = db.session.query(
        func.coalesce(func.sum(DeudaExtra.monto - DeudaExtra.pagado_monto), 0)).scalar()
    deuda_total = deuda_lecturas + deuda_extra

    residentes_en_mora = db.session.query(Residente.nombre,
                                           func.sum(LecturaAgua.total).label('deuda')) \
        .join(LecturaAgua).filter(LecturaAgua.pagado.is_(False)) \
        .group_by(Residente.id).order_by(func.sum(LecturaAgua.total).desc()).limit(8).all()

    ultimos_movimientos = MovimientoCaja.query.order_by(
        MovimientoCaja.fecha.desc(), MovimientoCaja.id.desc()).limit(8).all()

    return render_template('dashboard.html', cfg=cfg,
                            saldo_caja=saldo_caja, total_ingresos=total_ingresos,
                            total_egresos=total_egresos, ingresos_mes=ingresos_mes,
                            egresos_mes=egresos_mes, deuda_total=deuda_total,
                            residentes_en_mora=residentes_en_mora,
                            ultimos_movimientos=ultimos_movimientos)


# --------------------------------------------------------- RESIDENTES -----

@app.route('/residentes', methods=['GET', 'POST'])
@login_required
def residentes():
    if request.method == 'POST':
        r = Residente(nombre=request.form['nombre'].strip().upper(),
                      medidor=request.form.get('medidor', '').strip(),
                      telefono=request.form.get('telefono', '').strip())
        db.session.add(r)
        db.session.commit()
        flash('Vecino agregado.', 'success')
        return redirect(url_for('residentes'))
    lista = Residente.query.order_by(Residente.nombre).all()
    return render_template('residentes.html', residentes=lista)


@app.route('/residentes/<int:rid>/toggle')
@login_required
def toggle_residente(rid):
    r = Residente.query.get_or_404(rid)
    r.activo = not r.activo
    db.session.commit()
    return redirect(url_for('residentes'))


@app.route('/residentes/<int:rid>/eliminar')
@login_required
def eliminar_residente(rid):
    r = Residente.query.get_or_404(rid)
    db.session.delete(r)
    db.session.commit()
    flash('Vecino eliminado.', 'info')
    return redirect(url_for('residentes'))


# --------------------------------------------------------- MOVIMIENTOS ----

@app.route('/movimientos', methods=['GET', 'POST'])
@login_required
def movimientos():
    if request.method == 'POST':
        mov = MovimientoCaja(
            fecha=datetime.strptime(request.form['fecha'], '%Y-%m-%d').date(),
            tipo=request.form['tipo'],
            categoria=request.form.get('categoria', ''),
            detalle=request.form['detalle'].strip(),
            monto=float(request.form['monto']),
            medio_pago=request.form.get('medio_pago', 'Efectivo'),
            residente_id=request.form.get('residente_id') or None,
            observaciones=request.form.get('observaciones', ''))
        db.session.add(mov)
        db.session.commit()
        flash('Movimiento registrado.', 'success')
        return redirect(url_for('movimientos'))

    q = MovimientoCaja.query
    tipo = request.args.get('tipo')
    if tipo in ('INGRESO', 'EGRESO'):
        q = q.filter(MovimientoCaja.tipo == tipo)
    desde = request.args.get('desde')
    hasta = request.args.get('hasta')
    if desde:
        q = q.filter(MovimientoCaja.fecha >= desde)
    if hasta:
        q = q.filter(MovimientoCaja.fecha <= hasta)

    lista = q.order_by(MovimientoCaja.fecha.desc(), MovimientoCaja.id.desc()).all()
    residentes_lista = Residente.query.order_by(Residente.nombre).all()
    return render_template('movimientos.html', movimientos=lista,
                            residentes=residentes_lista, filtro_tipo=tipo,
                            desde=desde, hasta=hasta)


@app.route('/movimientos/<int:mid>/eliminar')
@login_required
def eliminar_movimiento(mid):
    m = MovimientoCaja.query.get_or_404(mid)
    db.session.delete(m)
    db.session.commit()
    flash('Movimiento eliminado.', 'info')
    return redirect(url_for('movimientos'))


# ------------------------------------------------------------ LECTURAS ----

@app.route('/lecturas', methods=['GET', 'POST'])
@login_required
def lecturas():
    cfg = get_config()
    if request.method == 'POST':
        residente_id = int(request.form['residente_id'])
        lect_ant = float(request.form['lectura_anterior'])
        lect_act = float(request.form['lectura_actual'])
        consumo = max(lect_act - lect_ant, 0)
        costo_agua = consumo * cfg.precio_m3
        monto_expensa = cfg.monto_expensa
        multa = float(request.form.get('multa') or 0)
        total = costo_agua + monto_expensa + multa

        lectura = LecturaAgua(
            residente_id=residente_id,
            mes=request.form['mes'],
            anio=int(request.form['anio']),
            fecha_lectura=datetime.strptime(request.form['fecha_lectura'], '%Y-%m-%d').date(),
            lectura_anterior=lect_ant,
            lectura_actual=lect_act,
            consumo_m3=consumo,
            costo_agua=costo_agua,
            monto_expensa=monto_expensa,
            multa=multa,
            total=total,
            observaciones=request.form.get('observaciones', ''))
        db.session.add(lectura)
        db.session.commit()
        flash(f'Lectura registrada. Total a cobrar: Bs {total:.2f}', 'success')
        return redirect(url_for('lecturas'))

    q = LecturaAgua.query
    residente_id = request.args.get('residente_id')
    mes = request.args.get('mes')
    anio = request.args.get('anio')
    if residente_id:
        q = q.filter(LecturaAgua.residente_id == int(residente_id))
    if mes:
        q = q.filter(LecturaAgua.mes == mes)
    if anio:
        q = q.filter(LecturaAgua.anio == int(anio))

    lista = q.order_by(LecturaAgua.anio.desc(), LecturaAgua.id.desc()).all()
    residentes_lista = Residente.query.filter_by(activo=True).order_by(Residente.nombre).all()
    return render_template('lecturas.html', lecturas=lista, residentes=residentes_lista,
                            meses=MESES, cfg=cfg, hoy=date.today())


@app.route('/lecturas/<int:lid>/pagar')
@login_required
def pagar_lectura(lid):
    l = LecturaAgua.query.get_or_404(lid)
    l.pagado = True
    l.fecha_pago = date.today()
    db.session.commit()
    flash('Pago registrado.', 'success')
    return redirect(url_for('lecturas'))


@app.route('/lecturas/<int:lid>/eliminar')
@login_required
def eliminar_lectura(lid):
    l = LecturaAgua.query.get_or_404(lid)
    db.session.delete(l)
    db.session.commit()
    return redirect(url_for('lecturas'))


# --------------------------------------------------------------- CUENTAS --

@app.route('/cuentas')
@login_required
def cuentas():
    residentes_lista = Residente.query.order_by(Residente.nombre).all()
    resumen = []
    for r in residentes_lista:
        pendientes = [l for l in r.lecturas if not l.pagado]
        deuda_agua_expensas = sum(l.total for l in pendientes)
        deuda_extra = sum(d.monto - d.pagado_monto for d in r.deudas_extra)
        resumen.append({
            'residente': r,
            'pendientes': pendientes,
            'deuda_agua_expensas': deuda_agua_expensas,
            'deuda_extra': deuda_extra,
            'deuda_total': deuda_agua_expensas + deuda_extra,
        })
    resumen.sort(key=lambda x: x['deuda_total'], reverse=True)
    return render_template('cuentas.html', resumen=resumen)


@app.route('/cuentas/deuda-extra', methods=['POST'])
@login_required
def agregar_deuda_extra():
    d = DeudaExtra(
        residente_id=request.form.get('residente_id') or None,
        nombre_deudor=request.form.get('nombre_deudor', ''),
        fecha=datetime.strptime(request.form['fecha'], '%Y-%m-%d').date(),
        concepto=request.form['concepto'],
        monto=float(request.form['monto']))
    db.session.add(d)
    db.session.commit()
    flash('Deuda registrada.', 'success')
    return redirect(url_for('cuentas'))


# ------------------------------------------------------------ CONFIG ------

@app.route('/config', methods=['GET', 'POST'])
@login_required
def config():
    cfg = get_config()
    if request.method == 'POST':
        cfg.nombre_edificio = request.form['nombre_edificio']
        cfg.monto_expensa = float(request.form['monto_expensa'])
        cfg.precio_m3 = float(request.form['precio_m3'])
        cfg.precio_limpieza_edificio = float(request.form['precio_limpieza_edificio'])
        cfg.precio_limpieza_parqueo = float(request.form['precio_limpieza_parqueo'])
        cfg.multa_mora = float(request.form.get('multa_mora') or 0)
        cfg.dia_limite_pago = int(request.form.get('dia_limite_pago') or 15)
        db.session.commit()
        flash('Configuración actualizada.', 'success')
        return redirect(url_for('config'))
    return render_template('config.html', cfg=cfg)


# ----------------------------------------------------------- REPORTES -----

@app.route('/reportes')
@login_required
def reportes():
    return render_template('reportes.html')


@app.route('/reportes/caja')
@login_required
def reporte_caja():
    desde = request.args.get('desde')
    hasta = request.args.get('hasta')
    q = MovimientoCaja.query
    if desde:
        q = q.filter(MovimientoCaja.fecha >= desde)
    if hasta:
        q = q.filter(MovimientoCaja.fecha <= hasta)
    movs = q.order_by(MovimientoCaja.fecha).all()
    total_ing = sum(m.monto for m in movs if m.tipo == 'INGRESO')
    total_egr = sum(m.monto for m in movs if m.tipo == 'EGRESO')
    return render_template('reporte_caja.html', movimientos=movs, total_ing=total_ing,
                            total_egr=total_egr, saldo=total_ing - total_egr,
                            desde=desde, hasta=hasta)


@app.route('/reportes/caja/excel')
@login_required
def reporte_caja_excel():
    desde = request.args.get('desde')
    hasta = request.args.get('hasta')
    q = MovimientoCaja.query
    if desde:
        q = q.filter(MovimientoCaja.fecha >= desde)
    if hasta:
        q = q.filter(MovimientoCaja.fecha <= hasta)
    movs = q.order_by(MovimientoCaja.fecha).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Libro Diario'
    headers = ['Fecha', 'Tipo', 'Detalle', 'Vecino', 'Medio de pago', 'Monto']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='4472C4')
    for m in movs:
        ws.append([m.fecha.strftime('%d/%m/%Y'), m.tipo, m.detalle,
                   m.residente.nombre if m.residente else '', m.medio_pago or '',
                   m.monto])
    total_ing = sum(m.monto for m in movs if m.tipo == 'INGRESO')
    total_egr = sum(m.monto for m in movs if m.tipo == 'EGRESO')
    ws.append([])
    ws.append(['', '', '', '', 'TOTAL INGRESOS', total_ing])
    ws.append(['', '', '', '', 'TOTAL EGRESOS', total_egr])
    ws.append(['', '', '', '', 'SALDO', total_ing - total_egr])
    for col, width in zip('ABCDEF', (12, 10, 45, 18, 14, 12)):
        ws.column_dimensions[col].width = width

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='libro_diario.xlsx',
                      mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/reportes/cuentas/excel')
@login_required
def reporte_cuentas_excel():
    residentes_lista = Residente.query.order_by(Residente.nombre).all()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Estado de Cuentas'
    ws.append(['Vecino', 'Deuda agua/expensas', 'Otras deudas', 'Deuda total'])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='4472C4')
    for r in residentes_lista:
        deuda_ae = sum(l.total for l in r.lecturas if not l.pagado)
        deuda_ex = sum(d.monto - d.pagado_monto for d in r.deudas_extra)
        ws.append([r.nombre, deuda_ae, deuda_ex, deuda_ae + deuda_ex])
    for col, width in zip('ABCD', (25, 20, 16, 16)):
        ws.column_dimensions[col].width = width

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='estado_de_cuentas.xlsx',
                      mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/reportes/lecturas/excel')
@login_required
def reporte_lecturas_excel():
    mes = request.args.get('mes')
    anio = request.args.get('anio')
    q = LecturaAgua.query
    if mes:
        q = q.filter(LecturaAgua.mes == mes)
    if anio:
        q = q.filter(LecturaAgua.anio == int(anio))
    lista = q.order_by(LecturaAgua.id).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Lecturas de agua'
    ws.append(['Vecino', 'Mes', 'Año', 'Lect. anterior', 'Lect. actual', 'Consumo m3',
               'Costo agua', 'Expensa', 'Multa', 'Total', 'Pagado'])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='4472C4')
    for l in lista:
        ws.append([l.residente.nombre, l.mes, l.anio, l.lectura_anterior, l.lectura_actual,
                   l.consumo_m3, l.costo_agua, l.monto_expensa, l.multa, l.total,
                   'SI' if l.pagado else 'NO'])
    for col, width in zip('ABCDEFGHIJK', (20, 12, 8, 14, 14, 12, 12, 10, 8, 10, 9)):
        ws.column_dimensions[col].width = width

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='lecturas_agua.xlsx',
                      mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ------------------------------------------------------------- CLI/INIT ---

@app.cli.command('init-db')
def init_db():
    """Crea las tablas y un usuario admin por defecto (flask init-db)."""
    db.create_all()
    if not Usuario.query.filter_by(username='admin').first():
        admin = Usuario(username='admin', nombre='Administrador',
                         password_hash=generate_password_hash('admin123'), rol='admin')
        db.session.add(admin)
    if not Configuracion.query.first():
        db.session.add(Configuracion(fecha_inicio_administracion=date.today()))
    db.session.commit()
    print('Base de datos inicializada. Usuario: admin / Clave: admin123 (¡cámbiala!)')


def ensure_bootstrap():
    """Crea tablas/usuario admin automáticamente si no existen (útil en el primer deploy)."""
    with app.app_context():
        db.create_all()
        if not Usuario.query.filter_by(username='admin').first():
            admin = Usuario(username='admin', nombre='Administrador',
                             password_hash=generate_password_hash('admin123'), rol='admin')
            db.session.add(admin)
        if not Configuracion.query.first():
            db.session.add(Configuracion(fecha_inicio_administracion=date.today()))
        db.session.commit()


ensure_bootstrap()

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=int(os.environ.get('PORT', 5000)))
